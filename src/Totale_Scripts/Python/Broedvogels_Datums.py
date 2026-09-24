# Let's recreate the variables and save the file properly since the execution state might have refreshed.
import pypdf
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import re

reader = pypdf.PdfReader("Handleiding-BMP-en-kolonievogels-2024-LR.pdf")

all_parsed_rows = []
for page_idx in range(42, 49):  # Pages 43 to 49
    page_text = reader.pages[page_idx].extract_text()
    lines = page_text.split('\n')
    for line in lines:
        line_clean = line.strip()
        if any(m in line_clean.lower() for m in
               ['jan', 'feb', 'mrt', 'apr', 'mei', 'jun', 'jul', 'aug', 'sep', 'okt', 'nov', 'dec']):
            if re.search(r'\d+\s+[a-z]{3}\s*-\s*\d*\s*[a-z]*', line_clean.lower()) or re.search(
                    r'\d+\s+[a-z]{3}\s*-\s*', line_clean.lower()):
                all_parsed_rows.append((page_idx + 1, line_clean))

parsed_records = []
for page_idx, line in all_parsed_rows:
    match = re.search(r'(\d+\s+[a-z]{3}\s*-\s*\d*\s*[a-z]{3})|(\d+\s+[a-z]{3}\s*-\s*)', line, re.IGNORECASE)
    if match:
        date_range = match.group(0).strip()
        before_date = line[:match.start()].strip()
        after_date = line[match.end():].strip()

        parts = before_date.split()
        name_parts = []
        col_parts = []
        for p in parts:
            if p.lower() in ['x', '.'] or p.isdigit():
                col_parts.append(p)
            else:
                name_parts.append(p)

        name = " ".join(name_parts)

        dist_match = re.search(r'\b\d+\b', after_date)
        distance = dist_match.group(0) if dist_match else ""

        parsed_records.append({
            "Pagina": page_idx,
            "Naam": name,
            "Kolom_Data": " ".join(col_parts),
            "Datumgrenzen": date_range,
            "Fusieafstand": distance,
            "Volledige_Lijn": line
        })

rows_to_save = []
for r in parsed_records:
    tokens = r['Kolom_Data'].split()
    while len(tokens) < 10:
        tokens.append("")

    rows_to_save.append({
        "Pagina": r['Pagina'],
        "Nederlandse Naam": r['Naam'],
        "BMP-Z": "✓" if (len(tokens) > 0 and tokens[0] == 'x') else "",
        "BMP-B": "✓" if (len(tokens) > 1 and tokens[1] == 'x') else "",
        "BMP-A": "✓" if (len(tokens) > 2 and tokens[2] == 'x') else "",
        "BMP-R": "✓" if (len(tokens) > 3 and tokens[3] == 'x') else "",
        "Kolonie": "✓" if (len(tokens) > 4 and tokens[4] == 'x') else "",
        "Datumgrenzen": r['Datumgrenzen'],
        "Fusieafstand (m)": r['Fusieafstand'],
        "Oorspronkelijke Lijn": r['Volledige_Lijn']
    })

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
ZEBRA_FILL = PatternFill(start_color="F2F6FA", end_color="F2F6FA", fill_type="solid")
WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FONT_TITLE = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
FONT_REGULAR = Font(name="Calibri", size=11, color="000000")
FONT_BOLD = Font(name="Calibri", size=11, bold=True, color="1F4E79")

BORDER_THIN = Border(
    left=Side(style='thin', color='D3D3D3'),
    right=Side(style='thin', color='D3D3D3'),
    top=Side(style='thin', color='D3D3D3'),
    bottom=Side(style='thin', color='D3D3D3')
)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Bijlage 2 Volledig"
ws.views.sheetView[0].showGridLines = True

ws.merge_cells("A1:J1")
title_cell = ws["A1"]
title_cell.value = "Bijlage 2: Volledige Soortenlijst per BMP-type en Interpretatiecriteria"
title_cell.font = FONT_TITLE
title_cell.fill = HEADER_FILL
title_cell.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 40

headers = ["Pagina", "Nederlandse Naam", "BMP-Z", "BMP-B", "BMP-A", "BMP-R", "Kolonievogel", "Datumgrenzen",
           "Fusieafstand (m)", "Oorspronkelijke Lijn in PDF"]
for col_num, h in enumerate(headers, 1):
    cell = ws.cell(row=3, column=col_num)
    cell.value = h
    cell.font = FONT_HEADER
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = BORDER_THIN
ws.row_dimensions[3].height = 25

for row_idx, r in enumerate(rows_to_save, 4):
    ws.row_dimensions[row_idx].height = 19
    is_zebra = (row_idx % 2 == 0)
    current_fill = ZEBRA_FILL if is_zebra else WHITE_FILL

    ws.cell(row=row_idx, column=1, value=r["Pagina"]).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=row_idx, column=2, value=r["Nederlandse Naam"]).font = FONT_BOLD

    for col_num, val in enumerate([r["BMP-Z"], r["BMP-B"], r["BMP-A"], r["BMP-R"], r["Kolonie"]], 3):
        cell_cb = ws.cell(row=row_idx, column=col_num, value=val)
        cell_cb.font = Font(name="Segoe UI Symbol", size=11, bold=True, color="1F4E79")
        cell_cb.alignment = Alignment(horizontal="center", vertical="center")

    ws.cell(row=row_idx, column=8, value=r["Datumgrenzen"]).alignment = Alignment(horizontal="center",
                                                                                  vertical="center")

    try:
        dist_val = int(r["Fusieafstand (m)"]) if r["Fusieafstand (m)"] else ""
    except ValueError:
        dist_val = r["Fusieafstand (m)"]
    c_dist = ws.cell(row=row_idx, column=9, value=dist_val)
    if isinstance(dist_val, int):
        c_dist.number_format = '#,##0'
    c_dist.alignment = Alignment(horizontal="right", vertical="center")

    ws.cell(row=row_idx, column=10, value=r["Oorspronkelijke Lijn"]).font = Font(name="Calibri", size=9, italic=True,
                                                                                 color="595959")

    # Apply global styles
    for col_num in range(1, 11):
        cell = ws.cell(row=row_idx, column=col_num)
        cell.fill = current_fill
        cell.border = BORDER_THIN
        if col_num not in [2, 10]:
            if col_num == 9:
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")

# Auto-adjust widths
for col in ws.columns:
    max_len = 0
    for cell in col:
        if cell.value:
            max_len = max(max_len, len(str(cell.value)))
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

ws.column_dimensions['A'].width = 10
ws.column_dimensions['B'].width = 28
ws.column_dimensions['H'].width = 18
ws.column_dimensions['I'].width = 18
ws.column_dimensions['J'].width = 50

ws.freeze_panes = "A4"

full_filename = "sovon_broedvogel_bijlage2_volledig_v2.xlsx"
wb.save(full_filename)
print(f"File saved: {full_filename}, {len(rows_to_save)} rows.")